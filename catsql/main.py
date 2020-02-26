#!/usr/bin/env python

from __future__ import print_function
from __future__ import unicode_literals

import argparse
import errno
from collections import OrderedDict
from datetime import datetime
from io import StringIO, BytesIO
import json
import openpyxl
import openpyxl.styles
import os
from shutil import copyfile
from sqlalchemy import Column, MetaData, Table, types
from sqlalchemy.exc import (CompileError, SAWarning)
from subprocess import call
import sys
import warnings

from catsql import cmdline
from catsql.database import Database
from catsql.nullify import Nullify
from catsql.patch import patchsql

if sys.version_info[0] == 2:
    import unicodecsv as csv
else:
    import csv

warnings.simplefilter("ignore", category=SAWarning)


def flatten_date(obj):
    if isinstance(obj, datetime):
        serial = obj.isoformat()
        return serial
    raise TypeError("No change needed")


# Get approximate length of header
class CsvRowWriter(object):
    def __init__(self):
        self.writer = None

    def writerow(self, row):
        row = [str(s) for s in row]
        try:
            queue = StringIO()
            self.writer = csv.writer(queue, lineterminator='')
            self.writer.writerow(row)
        except TypeError:
            queue = BytesIO()
            self.writer = csv.writer(queue, lineterminator='')
            self.writer.writerow(row)
        return queue.getvalue()


class SmartFormatter(argparse.HelpFormatter):

    def _split_lines(self, text, width):
        if text.startswith('R|'):
            return text[2:].splitlines()
        return argparse.HelpFormatter._split_lines(self, text, width)


class Viewer(object):
    def __init__(self, args, remainder, sys_args):
        self.args = args
        self.sys_args = sys_args
        self.failure = False
        self.values = None
        self.setup_filters(args)
        self.connect_database()
        self.process_remainder(remainder)

    def decomma(self, lst):
        if not lst:
            return lst
        lst = [l.split(',') for l in lst]
        return [item for sublist in lst for item in sublist]

    def setup_filters(self, args):
        self.url = args.catsql_database_url
        self.tables = self.decomma(args.table)
        self.schema = args.schema
        self.selected_columns = self.decomma(args.column)
        self.ordering = self.decomma(args.order)

        self.context_columns = set()
        self.context_filters = dict()
        if args.load_bookmark:
            with open(self.url, 'r') as fin:
                nargs = json.loads(fin.read())
            self.url = nargs['url']
            self.tables = set(nargs['table'])
            self.args.distinct = nargs.get('distinct', False)
            self.selected_columns = nargs.get('column', [])
            self.context_filters = nargs['context']
            self.context_columns = set(nargs['hidden_columns'])
            if args.value is None:
                self.args.value = []

        if self.tables is not None:
            self.tables = set(self.tables)
        self.row_filter = args.sql
        self.output_in_csv = args.csv
        self.output_in_json = args.json
        self.output_in_sqlite = args.sqlite
        self.output_in_excel = args.excel

        self.target_db = None
        if self.output_in_sqlite:
            self.target_db = Database(self.output_in_sqlite[0], can_create=True)
        self.target_ss = None
        if self.output_in_excel:
            self.target_ss = openpyxl.Workbook()
            self.target_ss.remove_sheet(self.target_ss.active)

        if args.value is not None:
            for context in args.value:
                if '=' in context:
                    key, value = context.split('=', 1)
                    self.context_filters[key] = value
                    self.context_columns.add(key)
                else:
                    self.context_columns.add(context)

    def connect_database(self):
        database = Database(self.url, verbose=self.args.verbose, tables=self.tables, schema=self.schema)
        self.database = database
        self.url = self.args.catsql_database_url = database.full_url

    def process_remainder(self, remainder):
        column_names = set()
        if len(remainder) == 0:
            return
        for table in self.database.tables_metadata.values():
            keys = None
            if hasattr(table, 'columns'):
                keys = table.columns.keys()
            else:
                keys = table.keys
            column_names |= set(keys)
        if 'catsql_database_url' in column_names:
            column_names.remove('catsql_database_url')
        parser = argparse.ArgumentParser()
        cmdline.add_options(parser)
        for name in sorted(column_names):
            try:
                parser.add_argument('--{}'.format(name), nargs=1, default=None)
            except argparse.ArgumentError:
                column_names.remove(name)
        self.values = vars(parser.parse_args(self.sys_args))
        dud = []
        for key, val in self.values.items():
            if (key not in column_names) or val is None:
                dud.append(key)
        for key in dud:
            del self.values[key]
        for key, val in self.values.items():
            self.values[key] = val[0]
            self.context_columns.add(key)

    def ok_column(self, name):
        if self.args.terse:
            if name in self.context_columns:
                return False
        if self.selected_columns:
            if name not in self.selected_columns:
                return False
        return True

    def start_table(self, table_name, columns):
        self.header_shown = False
        self.header_considered = False
        self.table_name = table_name
        self.columns = columns
        if self.selected_columns:
            self.columns = self.selected_columns

    def show_header_on_need(self):
        if self.header_shown or self.header_considered:
            return True
        self.header_considered = True
        if len(self.tables_so_far) > 0:
            if (self.output_in_csv or self.output_in_json or self.output_in_sqlite or
                  self.output_in_excel):
                if not self.output_in_sqlite:
                    self.failure = True
                self.tables_so_far.append(self.table_name)
                return False
            print("", file=self.output_file)
        if not (self.output_in_csv or self.output_in_json or self.output_in_sqlite or
                self.output_in_excel):
            print('== {} =='.format(self.table_name), file=self.output_file)
        if not (self.output_in_json or self.output_in_sqlite or self.output_in_excel):
            header_writer = CsvRowWriter()

            header = header_writer.writerow(list(column for column in self.columns
                                                 if self.ok_column(column)))
            print(header, file=self.output_file)
            if not self.output_in_csv:
                print('-' * len(header), file=self.output_file)
        self.header_shown = True
        self.tables_so_far.append(self.table_name)
        return True

    def show(self):

        self.tables_so_far = []

        work = None
        output_filename = None
        self.output_file = sys.stdout
        work_file = None
        self.start_table(None, None)

        try:
            if self.args.edit:
                import tempfile
                work = tempfile.mkdtemp()
                output_filename = os.path.join(work, 'reference.csv')
                work_file = open(output_filename, 'wt')
                self.output_file = work_file
                self.output_in_csv = True
                self.args.safe_null = True
                self.args.save_bookmark = [os.path.join(work, 'bookmark.json')]
            elif self.args.output:
                self.output_file = open(self.args.output[0], 'wt')

            viable_tables = []

            q = self.database.query(columns=self.selected_columns)
            touchable = True
            if self.args.select_from:
                q = q.select_from(self.args.select_from[0])
                touchable = False
            if self.args.distinct:
                q = q.distinct()
            if self.row_filter is not None:
                q = q.where_sqls(self.row_filter)
            if self.args.value is not None:
                q = q.where_kv(self.context_filters)
            if self.values is not None:
                q = q.where_kv_with_expansion(self.values)
            if self.args.grep:
                for pattern in self.args.grep:
                    q = q.grep(pattern, case_sensitive=False)
            if self.args.order:
                if 'none' not in self.args.order:
                    q.order(self.ordering)
            elif touchable:
                q.order()
            if self.args.limit:
                q = q.limit(int(self.args.limit[0]))

            ts = list(q)
            ts.sort(key=lambda t: t['table_name'])

            for t in ts:
                table_name = t['table_name']
                table = t['table']
                rows = t['rows']

                self.header_shown = False
                keys = None
                if hasattr(table, 'columns'):
                    keys = table.columns.keys()
                else:
                    keys = table.keys
                self.start_table(table_name, keys)
                viable_tables.append(table_name)

                if self.args.types:
                    column_types = []
                    for name in self.columns:
                        if not self.ok_column(name):
                            continue
                        try:
                            column = table.c[name]
                            sql_name = str(column.type)  # make sure not nulltype
                        except CompileError:
                            sql_name = None
                        column_types.append(sql_name)
                    rows = [column_types]

                if self.target_ss:
                    ws = self.target_ss.create_sheet()
                    ws.title = table_name
                    ws.append(column for column in self.columns
                                if self.ok_column(column))
                    for row in rows:
                        ws.append(cell for c, cell in enumerate(row)
                                if self.ok_column(self.columns[c]))
                    s = openpyxl.styles.NamedStyle(name="Header",
                                                   font=openpyxl.styles.Font(bold=True))
                    for cell in next(ws.rows):
                        cell.style = s
                    for column_cells in ws.columns:
                        ws.column_dimensions[column_cells[0].column_letter].auto_size = True

                if self.target_db:
                    if table_name in self.target_db.tables_metadata.keys():
                        # clear previous results
                        self.target_db.tables_metadata[table_name].drop(self.target_db.engine)
                    target = {'table': None, 'rows': []}

                    def fallback_type(example):
                        if isinstance(example, bool):
                            return types.Boolean
                        elif isinstance(example, int):
                            return types.Integer
                        elif isinstance(example, float):
                            return types.Float
                        elif isinstance(example, datetime):
                            return types.DateTime
                        return types.UnicodeText

                    def create_table(data):
                        if target['table'] is not None:
                            return
                        columns = []
                        for name in self.columns:
                            if not self.ok_column(name):
                                continue
                            column = table.c[name]
                            sql_type = column.type
                            try:
                                self.target_db.engine.dialect.type_compiler.process(sql_type)
                            except CompileError:
                                # some types need to be approximated
                                sql_type = None
                            if sql_type is None or isinstance(sql_type, types.NullType):
                                example = data.get(name)
                                sql_type = fallback_type(example)
                            sql_type.collation = None  # ignore collation
                            columns.append(Column(name, sql_type,
                                                  primary_key=column.primary_key))
                        metadata = MetaData(bind=self.target_db.engine)
                        target['table'] = Table(table_name, metadata, *columns)
                        target['table'].create(self.target_db.engine)

                    def add_row(data):
                        if data:
                            target['rows'].append(data)
                        if len(target['rows']) > 10000 or not data:
                            target['table'].insert().execute(target['rows'])
                            target['rows'] = []

                    def sqlited(data):
                        if isinstance(data, dict) or isinstance(data, list):
                            return json.dumps(data)
                        return data

                    for row in rows:
                        data = dict((self.columns[c], sqlited(cell))
                                    for c, cell in enumerate(row)
                                    if self.ok_column(self.columns[c]))
                        create_table(data)
                        add_row(data)
                    create_table({})
                    add_row(None)

                if self.output_in_json or self.output_in_sqlite or self.output_in_excel:
                    if not self.show_header_on_need():
                        continue
                    if self.output_in_json:
                        self.save_as_json(table, rows, self.output_in_json[0])
                elif not self.args.count:
                    # csv spec is that eol is \r\n; we ignore this for our purposes
                    # for good reasons that unfortunately there isn't space to describe
                    # here on the back of this envelope
                    csv_writer = csv.writer(self.output_file, lineterminator='\n')
                    if not self.show_header_on_need():
                        continue
                    if self.args.safe_null:
                        nullify = Nullify()
                        for row in rows:
                            csv_writer.writerow(list(nullify.encode_null(cell)
                                                     for c, cell in enumerate(row)
                                                     if self.ok_column(self.columns[c])))
                    else:
                        for row in rows:
                            csv_writer.writerow(list(cell for c, cell in enumerate(row)
                                                     if self.ok_column(self.columns[c])))
                    del csv_writer
                else:
                    self.show_header_on_need()
                    ct = rows.count()
                    print("({} row{})".format(ct, '' if ct == 1 else 's'), file=self.output_file)

            if len(self.tables_so_far) == 0 and len(viable_tables) == 1:
                self.show_header_on_need()

            if self.args.save_bookmark:
                with open(self.args.save_bookmark[0], 'w') as fout:
                    link = OrderedDict()
                    link['url'] = self.args.catsql_database_url
                    link['table'] = list(self.tables) if self.tables else None
                    link['column'] = self.selected_columns
                    link['distinct'] = self.args.distinct
                    link['context'] = self.context_filters
                    link['hidden_columns'] = sorted(self.context_columns)
                    link['sql'] = self.args.sql
                    fout.write(json.dumps(link, indent=2))

            if self.args.edit and not self.failure:
                work_file.close()
                work_file = None
                edit_filename = os.path.join(work, 'variant.csv')
                copyfile(output_filename, edit_filename)
                editor = os.environ.get('TABLE_EDITOR', None)
                if not editor:
                    editor = os.environ.get('EDITOR', 'nano')
                call([editor, edit_filename])
                patchsql([self.url, '--table'] + self.tables_so_far +
                         ['--follow', output_filename, edit_filename,
                            '--safe-null'] +
                            (['--quiet'] if self.args.quiet else []) +
                            (['--schema', self.schema] if self.schema else []),
                         database=self.database)

        finally:
            if self.failure:
                print("ERROR: "
                      "More than one table in csv/json output, consider adding:",
                      file=sys.stderr)
                for name in self.tables_so_far:
                    print("  --table {}".format(name),
                          file=sys.stderr)

            if self.target_ss:
                self.target_ss.save(self.output_in_excel[0])
            if work:
                if work_file:
                    try:
                        work_file.close()
                    except Exception:
                        pass
                    work_file = None
                import shutil
                shutil.rmtree(work)
                work = None

    def save_as_json(self, table, rows, filename):
        result = OrderedDict()
        result['count'] = 0
        result['meta'] = {
            'generator': 'catsql'
        }
        results = result['results'] = []
        names = []
        idxs = []
        for idx, column in enumerate(self.columns):
            if not self.ok_column(column):
                continue
            idxs.append(idx)
            names.append(column)
        for row in rows:
            od = OrderedDict()
            for i, idx in enumerate(idxs):
                od[names[i]] = row[idx]
            results.append(od)
        result['count'] = len(results)

        with open(filename, 'w') as fout:
            fout.write(json.dumps(result,
                                  indent=2,
                                  default=flatten_date))


def catsql(sys_args):

    parser = argparse.ArgumentParser(description='Quickly display and edit a slice of a database.',
                                     formatter_class=SmartFormatter)

    cmdline.add_options(parser)
    args, remainder = parser.parse_known_args(sys_args)
    viewer = Viewer(args, remainder, sys_args)
    viewer.show()


def main():
    try:
        catsql(sys.argv[1:])
    except IOError as error:
        if error.errno == errno.EPIPE:
            # totally benign e.g. pipe through head/tail
            pass
        else:
            raise


def edit(editor):
    args = sys.argv[1:]
    args += ['--edit']
    os.environ['TABLE_EDITOR'] = editor
    catsql(args)


if __name__ == "__main__":
    main()
